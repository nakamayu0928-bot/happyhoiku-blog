"""プレゼント用「すぐ使える経理・事務Excelテンプレ3点セット」を作るスクリプト。

使い方: python make_templates_xlsx.py
出力:   keiri-templates.xlsx（このファイルと同じフォルダ）

黄色のセル＝入力欄、白いセル＝自動計算。サンプル行は消して使う前提。
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("keiri-templates.xlsx")

FONT = "メイリオ"
ORANGE = "F28C28"
HEAD_FILL = PatternFill("solid", fgColor="FCE3C8")
INPUT_FILL = PatternFill("solid", fgColor="FFF6BF")
SUM_FILL = PatternFill("solid", fgColor="FFF3E6")
THIN = Side(style="thin", color="D9C2A7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
YEN = '"¥"#,##0;-"¥"#,##0;""'
DATE = "yyyy/mm/dd"
FIRST, LAST = 6, 205  # 入力行（200行分）


def f(size=10, bold=False, color="333333"):
    return Font(name=FONT, size=size, bold=bold, color=color)


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = f(14, True, ORANGE)
    ws["A2"] = sub
    ws["A2"].font = f(9, color="777777")
    ws["A3"] = "黄色＝入力するセル　白＝自動で計算（さわらない）　※サンプル行は消してから使ってください"
    ws["A3"].font = f(9, color="777777")


def header(ws, row, col, names, widths=None):
    for i, name in enumerate(names):
        c = ws.cell(row, col + i, name)
        c.font = f(10, True)
        c.fill = HEAD_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if widths:
            ws.column_dimensions[get_column_letter(col + i)].width = widths[i]


def styled(cell, fill=None, fmt=None, bold=False):
    cell.font = f(10, bold)
    cell.border = BORDER
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    return cell


def dropdown(ws, formula, rng):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error = "リストから選んでください（リストは「設定」シートで変えられます）"
    dv.errorTitle = "入力できない値です"
    ws.add_data_validation(dv)
    dv.add(rng)


def finish(ws, landscape=True):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


wb = Workbook()

# ================================================================ はじめに
ws = wb.active
ws.title = "はじめに"
ws["A1"] = "すぐ使える 経理・事務Excelテンプレ3点セット"
ws["A1"].font = f(16, True, ORANGE)
ws["A2"] = "作成: みかん🍊 @mikan_keiri"
ws["A2"].font = f(9, color="777777")
rows = [
    ("シート", "できること"),
    ("経費精算", "申請を入力するだけで、領収書もれ・二重申請・高額の申請に自動で印がつく。勘定科目別の合計も自動"),
    ("請求・入金", "請求と入金を入力するだけで「未入金」「期日超過」「手数料差引?」「一部入金」などの状況が自動で表示される"),
    ("月次集計", "経費（勘定科目別）と請求額（取引先別）を、月ごとに自動で集計。上司への報告にそのまま使える"),
    ("設定", "勘定科目・取引先・申請者のリストと、チェックの基準（金額・手数料）を自分の会社に合わせて変える場所"),
]
for r, (a, b) in enumerate(rows, 4):
    for c, v in ((1, a), (2, b)):
        cell = styled(ws.cell(r, c, v), HEAD_FILL if r == 4 else None, bold=(r == 4))
        cell.alignment = Alignment(wrap_text=True, vertical="center")
ws.column_dimensions["A"].width = 16
ws.column_dimensions["B"].width = 90

steps = [
    ("はじめの3ステップ", ""),
    ("", "① 「設定」シートで、勘定科目・取引先・申請者・年度の開始月を自分の会社に合わせる"),
    ("", "② 「経費精算」「請求・入金」のサンプル行（内容に「サンプル」と書いてある行）を消す"),
    ("", "③ 黄色のセルに入力していく。チェックや集計は自動で出ます"),
    ("色の意味", ""),
    ("", "黄色のセル … 入力するところ　／　白いセル … 自動で計算するところ（消すと計算されなくなります）"),
    ("", "オレンジ色の行 … チェックにひっかかった行（経費精算）、または期日超過・金額ちがい（請求・入金）"),
    ("自分用に直したいときは", ""),
    ("", "Claudeにこう聞くと、数式をこわさずに直す方法を教えてくれます。"),
    ("", "「このExcelテンプレの［経費精算］シートに［部門］の列を追加したいです。ほかのシートの数式がずれないように、手順を教えてください。"
         "シートの列は［A:No、B:申請日、C:申請者、D:勘定科目、E:内容、F:金額、G:領収書、H:承認、I:チェック］です」"),
    ("ご注意", ""),
    ("", "実際の業務で使う前に、会社のルール（経理規程・承認のながれ）と合っているか確認してください。"),
    ("", "テンプレのチェックは「見落とし防止」のためのものです。最終的な確認と判断は、担当者が行ってください。"),
]
for r, (a, b) in enumerate(steps, 11):
    ws.cell(r, 1, a).font = f(10, True, ORANGE)
    cell = ws.cell(r, 2, b)
    cell.font = f()
    cell.alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[20].height = 48
finish(ws, landscape=False)

# ================================================================ 設定
st = wb.create_sheet("設定")
title(st, "設定", "ここを変えると、ほかのシートのリストやチェックの基準が変わります。")
for r, (label, val, fmt, note) in enumerate([
    ("年度の開始年", 2026, "0", "例: 2026年4月始まりなら 2026"),
    ("年度の開始月", 4, "0", "例: 4月始まりなら 4、1月始まりなら 1"),
    ("要確認とする金額（経費）", 30000, YEN, "この金額以上の経費申請に「金額要確認」の印がつく"),
    ("手数料とみなす差額の上限（入金）", 880, YEN, "請求額より、この金額以内で少ない入金は「手数料差引?」と表示"),
], 5):
    styled(st.cell(r, 1, label), bold=True)
    styled(st.cell(r, 2, val), INPUT_FILL, fmt)
    st.cell(r, 3, note).font = f(9, color="777777")
st.column_dimensions["A"].width = 30
st.column_dimensions["B"].width = 14
st.column_dimensions["C"].width = 4

header(st, 10, 1, ["勘定科目（20まで）"])
kamoku = ["旅費交通費", "消耗品費", "会議費", "接待交際費", "通信費", "新聞図書費", "支払手数料", "福利厚生費", "雑費"]
for i in range(20):
    styled(st.cell(11 + i, 1, kamoku[i] if i < len(kamoku) else None), INPUT_FILL)
header(st, 10, 4, ["取引先（30まで）", "振込名義（カナ）"], [22, 22])
vendors = [("サンプル商事", "ｻﾝﾌﾟﾙｼﾖｳｼﾞ(ｶ"), ("ひだまり工業", "ﾋﾀﾞﾏﾘｺｳｷﾞﾖｳ(ｶ"), ("あおぞら設計", "ｱｵｿﾞﾗｾｯｹｲ(ﾕ")]
for i in range(30):
    v = vendors[i] if i < len(vendors) else (None, None)
    styled(st.cell(11 + i, 4, v[0]), INPUT_FILL)
    styled(st.cell(11 + i, 5, v[1]), INPUT_FILL)
header(st, 10, 7, ["申請者（30まで）"], [18])
for i, name in enumerate(["佐藤", "鈴木"] + [None] * 28):
    styled(st.cell(11 + i, 7, name), INPUT_FILL)
st.column_dimensions["F"].width = 4
finish(st, landscape=False)

KAMOKU = "'設定'!$A$11:$A$30"
VENDOR = "'設定'!$D$11:$D$40"
STAFF = "'設定'!$G$11:$G$40"

# ================================================================ 経費精算
ws = wb.create_sheet("経費精算")
title(ws, "経費精算一覧", "申請を1行ずつ入力します。I列のチェックが空なら問題なしです。")
header(ws, 5, 1, ["No", "申請日", "申請者", "勘定科目", "内容", "金額", "領収書", "承認", "チェック"],
       [6, 12, 11, 13, 28, 12, 8, 8, 24])
samples = [
    (date(2026, 10, 1), "佐藤", "旅費交通費", "客先訪問 電車代（サンプル）", 640, "あり", "済"),
    (date(2026, 10, 2), "鈴木", "消耗品費", "コピー用紙（サンプル）", 1980, "なし", "未"),
    (date(2026, 10, 3), "佐藤", "会議費", "打ち合わせ 会場代（サンプル）", 35000, "あり", "未"),
]
for r in range(FIRST, LAST + 1):
    styled(ws.cell(r, 1, f'=IF(B{r}="","",ROW()-{FIRST - 1})'))
    for c, fmt in ((2, DATE), (3, None), (4, None), (5, None), (6, YEN), (7, None), (8, None)):
        styled(ws.cell(r, c), INPUT_FILL, fmt)
    styled(ws.cell(r, 9, (
        f'=IF(B{r}="","",TRIM(IF(G{r}="なし","領収書なし ","")'
        f'&IF(COUNTIFS($B${FIRST}:$B${LAST},B{r},$C${FIRST}:$C${LAST},C{r},$F${FIRST}:$F${LAST},F{r})>1,"重複? ","")'
        f"&IF(F{r}>='設定'!$B$7,\"金額要確認\",\"\")))"
    )))
for i, row in enumerate(samples):
    for c, v in enumerate(row, 2):
        ws.cell(FIRST + i, c, v)
dropdown(ws, STAFF, f"C{FIRST}:C{LAST}")
dropdown(ws, KAMOKU, f"D{FIRST}:D{LAST}")
dropdown(ws, '"あり,なし"', f"G{FIRST}:G{LAST}")
dropdown(ws, '"未,済"', f"H{FIRST}:H{LAST}")
ws.conditional_formatting.add(
    f"A{FIRST}:I{LAST}",
    FormulaRule(formula=[f'$I{FIRST}<>""'], fill=PatternFill("solid", fgColor="FDDCB5")))
ws.freeze_panes = f"A{FIRST}"

# 集計（右側）
ws["K4"] = "勘定科目別の合計"
ws["K4"].font = f(10, True, ORANGE)
header(ws, 5, 11, ["勘定科目", "合計", "件数"], [14, 12, 7])
for i in range(20):
    r = 6 + i
    styled(ws.cell(r, 11, f"=IF('設定'!A{11 + i}=\"\",\"\",'設定'!A{11 + i})"))
    styled(ws.cell(r, 12, f'=IF(K{r}="","",SUMIFS($F${FIRST}:$F${LAST},$D${FIRST}:$D${LAST},K{r}))'), fmt=YEN)
    styled(ws.cell(r, 13, f'=IF(K{r}="","",COUNTIFS($D${FIRST}:$D${LAST},K{r}))'))
styled(ws.cell(26, 11, "合計"), SUM_FILL, bold=True)
styled(ws.cell(26, 12, f"=SUM(F{FIRST}:F{LAST})"), SUM_FILL, YEN, True)
styled(ws.cell(26, 13, f"=COUNT(F{FIRST}:F{LAST})"), SUM_FILL, bold=True)
for r, (label, formula) in enumerate([
    ("チェックがついた件数", f'=COUNTIF(I{FIRST}:I{LAST},"?*")'),
    ("承認がまだの件数", f'=COUNTIFS(B{FIRST}:B{LAST},"<>",H{FIRST}:H{LAST},"<>済")'),
], 28):
    styled(ws.cell(r, 11, label), bold=True)
    styled(ws.cell(r, 12, formula))
ws.column_dimensions["J"].width = 3
finish(ws)

# ================================================================ 請求・入金
ws = wb.create_sheet("請求・入金")
title(ws, "請求・入金管理表", "請求を出したら左側を、入金があったら入金日と入金額を入力します。状況は自動で表示されます。")
header(ws, 5, 1, ["請求No", "請求日", "取引先", "請求額", "支払期日", "入金日", "入金額", "差額", "状況", "メモ"],
       [11, 12, 16, 12, 12, 12, 12, 10, 13, 24])
samples = [
    ("INV-001", date(2026, 9, 30), "サンプル商事", 110000, date(2026, 10, 31), date(2026, 10, 25), 109340, "（サンプル）"),
    ("INV-002", date(2026, 9, 30), "ひだまり工業", 55000, date(2026, 10, 31), None, None, "（サンプル）"),
    ("INV-003", date(2026, 9, 30), "あおぞら設計", 33000, date(2026, 10, 31), date(2026, 10, 30), 33000, "（サンプル）"),
]
for r in range(FIRST, LAST + 1):
    for c, fmt in ((1, None), (2, DATE), (3, None), (4, YEN), (5, DATE), (6, DATE), (7, YEN), (10, None)):
        styled(ws.cell(r, c), INPUT_FILL, fmt)
    styled(ws.cell(r, 8, f'=IF(OR(D{r}="",G{r}=""),"",G{r}-D{r})'), fmt=YEN)
    styled(ws.cell(r, 9, (
        f'=IF(D{r}="","",IF(G{r}="",IF(AND(E{r}<>"",TODAY()>E{r}),"期日超過","未入金"),'
        f'IF(G{r}=D{r},"入金済",IF(AND(D{r}-G{r}>0,D{r}-G{r}<=\'設定\'!$B$8),"手数料差引?",'
        f'IF(G{r}<D{r},"一部入金","過入金")))))'
    )))
for i, row in enumerate(samples):
    for c, v in zip((1, 2, 3, 4, 5, 6, 7, 10), row):
        ws.cell(FIRST + i, c, v)
dropdown(ws, VENDOR, f"C{FIRST}:C{LAST}")
for text, color in (("期日超過", "F8C9C9"), ("一部入金", "FDDCB5"), ("過入金", "FDDCB5"), ("手数料差引?", "FFF0B3"), ("入金済", "DDEFD6")):
    ws.conditional_formatting.add(
        f"I{FIRST}:I{LAST}",
        FormulaRule(formula=[f'$I{FIRST}="{text}"'], fill=PatternFill("solid", fgColor=color)))
ws.freeze_panes = f"A{FIRST}"

ws["L4"] = "状況のまとめ"
ws["L4"].font = f(10, True, ORANGE)
header(ws, 5, 12, ["状況", "件数", "請求額の合計"], [14, 7, 14])
for i, s in enumerate(["未入金", "期日超過", "一部入金", "手数料差引?", "過入金", "入金済"]):
    r = 6 + i
    styled(ws.cell(r, 12, s))
    styled(ws.cell(r, 13, f'=COUNTIF($I${FIRST}:$I${LAST},L{r})'))
    styled(ws.cell(r, 14, f'=SUMIFS($D${FIRST}:$D${LAST},$I${FIRST}:$I${LAST},L{r})'), fmt=YEN)
styled(ws.cell(12, 12, "請求額−入金額（差し引き）"), SUM_FILL, bold=True)
styled(ws.cell(12, 13, "", ), SUM_FILL)
styled(ws.cell(12, 14, f'=SUM(D{FIRST}:D{LAST})-SUM(G{FIRST}:G{LAST})'), SUM_FILL, YEN, True)
ws["L14"] = "「手数料差引?」は振込手数料を差し引かれた可能性があります。"
ws["L15"] = "取引条件（手数料はどちらの負担か）を確認してから処理してください。"
for c in ("L14", "L15"):
    ws[c].font = f(9, color="777777")
ws.column_dimensions["K"].width = 3
finish(ws)

# ================================================================ 月次集計
ws = wb.create_sheet("月次集計")
title(ws, "月次集計", "「経費精算」「請求・入金」に入力した内容を、月ごとに自動で集計します。年度は「設定」シートで変えられます。")
ws["A3"] = "このシートは入力不要です（すべて自動で計算されます）"
MONTH_FMT = 'yyyy"年"m"月"'


def month_header(row, label):
    styled(ws.cell(row, 1, label), HEAD_FILL, bold=True)
    for i in range(12):
        col = 2 + i
        formula = "=DATE('設定'!$B$5,'設定'!$B$6,1)" if (i == 0 and row == 5) else (
            f"=EDATE({get_column_letter(col - 1)}{row},1)" if row == 5 else f"={get_column_letter(col)}$5")
        c = styled(ws.cell(row, col, formula), HEAD_FILL, MONTH_FMT, True)
        c.alignment = Alignment(horizontal="center")
    styled(ws.cell(row, 14, "年度合計"), HEAD_FILL, bold=True)


ws["A4"] = "経費（勘定科目別）"
ws["A4"].font = f(11, True, ORANGE)
month_header(5, "勘定科目")
for i in range(20):
    r = 6 + i
    styled(ws.cell(r, 1, f"=IF('設定'!A{11 + i}=\"\",\"\",'設定'!A{11 + i})"))
    for m in range(12):
        col = get_column_letter(2 + m)
        styled(ws.cell(r, 2 + m, (
            f"=IF($A{r}=\"\",\"\",SUMIFS('経費精算'!$F${FIRST}:$F${LAST},'経費精算'!$D${FIRST}:$D${LAST},$A{r},"
            f"'経費精算'!$B${FIRST}:$B${LAST},\">=\"&{col}$5,'経費精算'!$B${FIRST}:$B${LAST},\"<\"&EDATE({col}$5,1)))"
        )), fmt=YEN)
    styled(ws.cell(r, 14, f'=IF($A{r}="","",SUM(B{r}:M{r}))'), fmt=YEN, bold=True)
styled(ws.cell(26, 1, "経費の合計"), SUM_FILL, bold=True)
for c in range(2, 15):
    L = get_column_letter(c)
    styled(ws.cell(26, c, f"=SUM({L}6:{L}25)"), SUM_FILL, YEN, True)

ws["A29"] = "請求額（取引先別・請求日の月で集計）"
ws["A29"].font = f(11, True, ORANGE)
month_header(30, "取引先")
for i in range(30):
    r = 31 + i
    styled(ws.cell(r, 1, f"=IF('設定'!D{11 + i}=\"\",\"\",'設定'!D{11 + i})"))
    for m in range(12):
        col = get_column_letter(2 + m)
        styled(ws.cell(r, 2 + m, (
            f"=IF($A{r}=\"\",\"\",SUMIFS('請求・入金'!$D${FIRST}:$D${LAST},'請求・入金'!$C${FIRST}:$C${LAST},$A{r},"
            f"'請求・入金'!$B${FIRST}:$B${LAST},\">=\"&{col}$5,'請求・入金'!$B${FIRST}:$B${LAST},\"<\"&EDATE({col}$5,1)))"
        )), fmt=YEN)
    styled(ws.cell(r, 14, f'=IF($A{r}="","",SUM(B{r}:M{r}))'), fmt=YEN, bold=True)
styled(ws.cell(61, 1, "請求額の合計"), SUM_FILL, bold=True)
for c in range(2, 15):
    L = get_column_letter(c)
    styled(ws.cell(61, c, f"=SUM({L}31:{L}60)"), SUM_FILL, YEN, True)
styled(ws.cell(62, 1, "入金額の合計（入金日の月）"), SUM_FILL, bold=True)
for m in range(12):
    col = get_column_letter(2 + m)
    styled(ws.cell(62, 2 + m, (
        f"=SUMIFS('請求・入金'!$G${FIRST}:$G${LAST},'請求・入金'!$F${FIRST}:$F${LAST},\">=\"&{col}$5,"
        f"'請求・入金'!$F${FIRST}:$F${LAST},\"<\"&EDATE({col}$5,1))"
    )), SUM_FILL, YEN, True)
styled(ws.cell(62, 14, "=SUM(B62:M62)"), SUM_FILL, YEN, True)
ws.column_dimensions["A"].width = 26
for c in range(2, 15):
    ws.column_dimensions[get_column_letter(c)].width = 11
ws.freeze_panes = "B6"
finish(ws)

for sheet in wb.worksheets:
    sheet.sheet_properties.tabColor = ORANGE if sheet.title in ("はじめに", "設定") else "F7C08A"
wb.move_sheet("設定", offset=3)  # はじめに → 経費精算 → 請求・入金 → 月次集計 → 設定
wb.save(OUT)
print(OUT)
