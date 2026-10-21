"""プレゼント用「練習用シート」(Excel) を作るスクリプト。

使い方: python make_practice_xlsx.py
出力:   claude-excel-practice.xlsx（このファイルと同じフォルダ）

データはすべて架空のダミーです。PDF のプロンプトで練習できるよう、
表記ゆれ・重複・入金額のズレなどを「わざと」入れています。
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).with_name("claude-excel-practice.xlsx")

FONT = "メイリオ"
ORANGE = "F28C28"
HEAD_FILL = PatternFill("solid", fgColor="FCE3C8")
INPUT_FILL = PatternFill("solid", fgColor="FFF6BF")
TRAP_FILL = PatternFill("solid", fgColor="FDE2E2")
THIN = Side(style="thin", color="D9C2A7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

YEN = '"¥"#,##0'


def f(size=10, bold=False, color="333333"):
    return Font(name=FONT, size=size, bold=bold, color=color)


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = f(14, True, ORANGE)
    ws["A2"] = sub
    ws["A2"].font = f(9, color="777777")


def table(ws, top, headers, rows, widths, formats=None):
    """top 行目に見出し、その下に rows を書く。数式は文字列のまま渡す。"""
    formats = formats or {}
    for c, h in enumerate(headers, 1):
        cell = ws.cell(top, c, h)
        cell.font = f(10, True)
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r, row in enumerate(rows, top + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.font = f()
            cell.border = BORDER
            if c in formats:
                cell.number_format = formats[c]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(top + 1, 1)


def mark(cell, fill, note):
    cell.fill = fill
    cell.comment = Comment(note, "みかん")


wb = Workbook()

# ---------------------------------------------------------------- はじめに
ws = wb.active
ws.title = "はじめに"
title(ws, "Claude×Excel 練習用シート", "プロンプト集（PDF）とセットで使ってください。データはすべて架空のダミーです。")
guide = [
    ("シート", "練習できること", "対応するプロンプト（PDF）"),
    ("売上表", "月別・取引先別の集計、商品マスタからの単価の取り出し、合計が合わない原因さがし", "No.1〜3、6〜8"),
    ("商品マスタ", "VLOOKUP / XLOOKUP で売上表とつなぐ", "No.3、6"),
    ("経費精算", "科目別の集計、重複・異常値・領収書もれのチェック", "No.2、15、16"),
    ("取引先リスト", "表記ゆれ（株式会社／(株)／全角半角など）の統一、重複チェック", "No.10〜12"),
    ("請求と入金", "請求一覧と入金一覧の照合（入金もれ・金額ちがい）", "No.14"),
]
for r, row in enumerate(guide, 4):
    for c, v in enumerate(row, 1):
        cell = ws.cell(r, c, v)
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.font = f(10, r == 4)
        if r == 4:
            cell.fill = HEAD_FILL
ws.column_dimensions["A"].width = 16
ws.column_dimensions["B"].width = 62
ws.column_dimensions["C"].width = 26

notes = [
    ("色の意味", ""),
    ("", "黄色のセル … 数式を入れて練習する場所です。自由に書きかえてOK。"),
    ("", "ピンクのセル … わざと入れた「ミス」です。カーソルを合わせるとヒントが出ます。"),
    ("使い方", ""),
    ("", "① PDFのプロンプトをコピーして、［ ］の部分をこのシートの内容に書きかえる"),
    ("", "② シートの表をClaudeに貼り付ける（またはファイルごと添付する）"),
    ("", "③ 返ってきた数式を黄色のセルに入れて、結果を確かめる"),
    ("注意", ""),
    ("", "実際の業務データを使うときは、会社のルールを確認し、取引先名・個人名などを伏せてから使ってください。"),
]
for r, (a, b) in enumerate(notes, 12):
    ws.cell(r, 1, a).font = f(10, True, ORANGE)
    ws.cell(r, 2, b).font = f()
ws["A23"] = "作成: @mikan_keiri"
ws["A23"].font = f(9, color="777777")

# ---------------------------------------------------------------- 商品マスタ
ws = wb.create_sheet("商品マスタ")
title(ws, "商品マスタ", "売上表の「商品コード」から商品名・単価を取り出す練習用です。")
products = [
    ("A-101", "A4コピー用紙（500枚×5冊）", 2800),
    ("A-102", "ボールペン 黒（10本）", 980),
    ("A-103", "クリアファイル（100枚）", 1500),
    ("B-201", "トナーカートリッジ", 12800),
    ("B-202", "ラベルシール（20シート）", 1200),
    ("C-301", "デスクライト", 6400),
]
table(ws, 4, ["商品コード", "商品名", "単価"], products, [12, 30, 12], {3: YEN})

# ---------------------------------------------------------------- 売上表
ws = wb.create_sheet("売上表")
title(ws, "売上表（2026年8月〜10月）", "月別・取引先別の集計や、商品マスタとの突き合わせの練習用です。")
clients = ["ひだまり工業", "サンプル商事", "あおぞら設計", "こもれび食品"]
plan = [  # (月, 日, 取引先index, 商品コード, 数量)
    (8, 3, 0, "A-101", 4), (8, 5, 1, "A-102", 10), (8, 8, 2, "B-201", 2),
    (8, 12, 3, "A-103", 3), (8, 18, 0, "B-202", 5), (8, 22, 1, "C-301", 1),
    (8, 27, 2, "A-101", 6), (8, 30, 3, "A-102", 8), (9, 2, 0, "B-201", 1),
    (9, 6, 1, "A-101", 5), (9, 10, 2, "A-103", 2), (9, 13, 3, "B-202", 4),
    (9, 17, 0, "A-102", 12), (9, 21, 1, "B-201", 3), (9, 24, 2, "C-301", 2),
    (9, 28, 3, "A-101", 7), (10, 1, 0, "A-103", 6), (10, 4, 1, "B-202", 2),
    (10, 8, 2, "A-102", 15), (10, 11, 3, "B-201", 2), (10, 15, 0, "C-301", 1),
    (10, 19, 1, "A-101", 3), (10, 23, 2, "B-202", 6), (10, 27, 3, "A-103", 5),
]
rows = []
for i, (m, d, ci, code, qty) in enumerate(plan):
    r = 5 + i
    rows.append([
        date(2026, m, d), f"S-{2026}{m:02d}{d:02d}", clients[ci], code, qty,
        f"=INDEX(商品マスタ!$C$5:$C$10,MATCH(D{r},商品マスタ!$A$5:$A$10,0))",
        f"=E{r}*F{r}",
    ])
table(ws, 4, ["日付", "伝票No", "取引先", "商品コード", "数量", "単価", "金額"], rows,
      [12, 15, 16, 12, 8, 11, 12], {1: "yyyy/mm/dd", 6: YEN, 7: YEN})
last = 4 + len(rows)
ws.cell(last + 1, 6, "合計").font = f(10, True)
ws.cell(last + 1, 7, f"=SUM(G5:G{last})").number_format = YEN
ws.cell(last + 1, 7).font = f(10, True)
ws.cell(last + 1, 4, "数量合計").font = f(10, True)
ws.cell(last + 1, 5, f"=SUM(E5:E{last})").font = f(10, True)

# わざと入れたミス: 数量が「文字列の数字」になっていて SUM に入らない
trap = ws.cell(13, 5)  # 9/2 の行
trap.value = "1"
mark(trap, TRAP_FILL, "この「1」は文字として入っています。\n数量合計(SUM)に入らないので、合計が合いません。\nPDFのNo.8で原因をClaudeに聞いてみましょう。")
trap.alignment = Alignment(horizontal="left")

# 練習欄: 月別・取引先別の集計
ws["J4"] = "練習① 月別の売上合計（PDF No.1）"
ws["J4"].font = f(10, True, ORANGE)
for i, (label, ym) in enumerate([("2026年8月", "2026/8/1"), ("2026年9月", "2026/9/1"), ("2026年10月", "2026/10/1")]):
    ws.cell(5 + i, 10, label).font = f()
    c = ws.cell(5 + i, 11)
    c.fill = INPUT_FILL
    c.border = BORDER
    c.number_format = YEN
ws["J9"] = "練習② 取引先別の売上合計（PDF No.2）"
ws["J9"].font = f(10, True, ORANGE)
for i, name in enumerate(clients):
    ws.cell(10 + i, 10, name).font = f()
    c = ws.cell(10 + i, 11)
    c.fill = INPUT_FILL
    c.border = BORDER
    c.number_format = YEN
ws["J15"] = "黄色のセルに、Claudeに作ってもらった数式を入れてみましょう。"
ws["J15"].font = f(9, color="777777")
ws["J16"] = "答え合わせ: 3か月の合計は G列の一番下の「合計」と同じになります。"
ws["J16"].font = f(9, color="777777")
ws.column_dimensions["J"].width = 20
ws.column_dimensions["K"].width = 14

# ---------------------------------------------------------------- 経費精算
ws = wb.create_sheet("経費精算")
title(ws, "経費精算一覧（2026年10月分）", "科目別の集計や、申請内容のチェックの練習用です。")
exp = [
    (date(2026, 10, 1), "佐藤", "旅費交通費", "客先訪問 電車代", 640, "あり"),
    (date(2026, 10, 2), "鈴木", "消耗品費", "プリンター用紙", 1980, "あり"),
    (date(2026, 10, 3), "高橋", "会議費", "打ち合わせ コーヒー代", 1320, "あり"),
    (date(2026, 10, 6), "佐藤", "旅費交通費", "客先訪問 バス代", 460, "あり"),
    (date(2026, 10, 7), "田中", "通信費", "切手代", 840, "あり"),
    (date(2026, 10, 8), "鈴木", "旅費交通費", "研修会場 電車代", 720, "あり"),
    (date(2026, 10, 9), "高橋", "消耗品費", "USBメモリ", 1580, "なし"),
    (date(2026, 10, 10), "田中", "旅費交通費", "客先訪問 電車代", 48000, "あり"),
    (date(2026, 10, 14), "佐藤", "会議費", "来客用お茶菓子", 2160, "あり"),
    (date(2026, 10, 15), "鈴木", "旅費交通費", "研修会場 電車代", 720, "あり"),
    (date(2026, 10, 15), "鈴木", "旅費交通費", "研修会場 電車代", 720, "あり"),
    (date(2026, 10, 16), "高橋", "通信費", "レターパック", 1200, "あり"),
    (date(2026, 10, 20), "田中", "消耗品費", "電池", 650, "あり"),
    (date(2026, 10, 21), "佐藤", "旅費交通費", "客先訪問 タクシー代", 2380, "あり"),
    (date(2026, 10, 23), "鈴木", "会議費", "打ち合わせ 昼食代", 3300, "あり"),
    (date(2026, 10, 27), "高橋", "旅費交通費", "客先訪問 電車代", 580, "あり"),
]
table(ws, 4, ["申請日", "申請者", "勘定科目", "内容", "金額", "領収書"], exp,
      [12, 10, 13, 24, 11, 9], {1: "yyyy/mm/dd", 5: YEN})
last = 4 + len(exp)
ws.cell(last + 1, 4, "合計").font = f(10, True)
ws.cell(last + 1, 5, f"=SUM(E5:E{last})").number_format = YEN
ws.cell(last + 1, 5).font = f(10, True)
mark(ws["E12"], TRAP_FILL, "電車代にしては高すぎる金額です（桁の打ちまちがい？）。\nPDFのNo.16で見つけ方を練習できます。")
mark(ws["F11"], TRAP_FILL, "領収書が「なし」です。")
mark(ws["A15"], TRAP_FILL, "上の行とまったく同じ内容です（二重申請？）。")

ws["H4"] = "練習③ 勘定科目別の合計（PDF No.2）"
ws["H4"].font = f(10, True, ORANGE)
for i, k in enumerate(["旅費交通費", "消耗品費", "会議費", "通信費"]):
    ws.cell(5 + i, 8, k).font = f()
    c = ws.cell(5 + i, 9)
    c.fill = INPUT_FILL
    c.border = BORDER
    c.number_format = YEN
ws.column_dimensions["H"].width = 14
ws.column_dimensions["I"].width = 13

# ---------------------------------------------------------------- 取引先リスト
ws = wb.create_sheet("取引先リスト")
title(ws, "取引先リスト（表記ゆれあり）", "同じ会社がちがう書き方で登録されています。統一と重複チェックの練習用です。")
vendors = [
    ("V001", "株式会社ひだまり工業", "03-1234-5678"),
    ("V002", "(株)ひだまり工業", "03-1234-5678"),
    ("V003", "サンプル商事株式会社", "06-2345-6789"),
    ("V004", "ｻﾝﾌﾟﾙ商事（株）", "06-2345-6789"),
    ("V005", "有限会社あおぞら設計", "045-345-6789"),
    ("V006", "あおぞら設計　", "045-345-6789"),
    ("V007", "株式会社こもれび食品", "052-456-7890"),
    ("V008", "株式会社 こもれび食品", "052-456-7890"),
    ("V009", "合同会社ひなた印刷", "011-567-8901"),
    ("V010", "ＡＢＣ物流株式会社", "092-678-9012"),
    ("V011", "ABC物流（株）", "092-678-9012"),
    ("V012", "株式会社みどり不動産", "03-7890-1234"),
]
table(ws, 4, ["取引先ID", "取引先名", "電話番号"], vendors, [10, 28, 16])
ws["E4"] = "統一後の名前（練習④ PDF No.10・11）"
ws["E4"].font = f(10, True, ORANGE)
for r in range(5, 5 + len(vendors)):
    c = ws.cell(r, 5)
    c.fill = INPUT_FILL
    c.border = BORDER
ws.column_dimensions["E"].width = 30
ws["E18"] = "ヒント: 電話番号が同じなら、同じ会社の可能性が高いです。"
ws["E18"].font = f(9, color="777777")

# ---------------------------------------------------------------- 請求と入金
ws = wb.create_sheet("請求と入金")
title(ws, "請求と入金の照合（2026年9月請求分）", "左が請求一覧、右が通帳の入金一覧です。入金もれ・金額ちがいを探す練習用です。")
bills = [
    ("INV-0901", "ひだまり工業", 132000, date(2026, 10, 31)),
    ("INV-0902", "サンプル商事", 85800, date(2026, 10, 31)),
    ("INV-0903", "あおぞら設計", 47300, date(2026, 10, 31)),
    ("INV-0904", "こもれび食品", 210100, date(2026, 10, 31)),
    ("INV-0905", "ひなた印刷", 19800, date(2026, 10, 31)),
]
table(ws, 4, ["請求No", "取引先", "請求額", "支払期日"], bills, [12, 15, 12, 12], {3: YEN, 4: "yyyy/mm/dd"})
pays = [
    (date(2026, 10, 24), "ﾋﾀﾞﾏﾘｺｳｷﾞﾖｳ(ｶ", 132000),
    (date(2026, 10, 28), "ｻﾝﾌﾟﾙｼﾖｳｼﾞ(ｶ", 85140),
    (date(2026, 10, 30), "ｺﾓﾚﾋﾞｼﾖｸﾋﾝ(ｶ", 210100),
    (date(2026, 10, 30), "ﾋﾅﾀｲﾝｻﾂ(ﾄﾞ", 19800),
]
for c, h in enumerate(["入金日", "振込名義", "入金額"], 6):
    cell = ws.cell(4, c, h)
    cell.font = f(10, True)
    cell.fill = HEAD_FILL
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="center")
for r, row in enumerate(pays, 5):
    for c, v in enumerate(row, 6):
        cell = ws.cell(r, c, v)
        cell.font = f()
        cell.border = BORDER
        if c == 6:
            cell.number_format = "yyyy/mm/dd"
        if c == 8:
            cell.number_format = YEN
ws.column_dimensions["E"].width = 3
ws.column_dimensions["F"].width = 12
ws.column_dimensions["G"].width = 20
ws.column_dimensions["H"].width = 12
mark(ws["H6"], TRAP_FILL, "請求額より660円少ない入金です。\n振込手数料を差し引かれた可能性があります。")
ws["A11"] = "照合結果（練習⑤ PDF No.14）"
ws["A11"].font = f(10, True, ORANGE)
for c, h in enumerate(["請求No", "入金状況", "差額", "メモ"], 1):
    cell = ws.cell(12, c, h)
    cell.font = f(10, True)
    cell.fill = HEAD_FILL
    cell.border = BORDER
for r in range(13, 18):
    for c in range(1, 5):
        cell = ws.cell(r, c)
        cell.fill = INPUT_FILL
        cell.border = BORDER
ws["A19"] = "ヒント: 入金一覧に名前がない会社と、金額が合わない会社がそれぞれ1社あります。"
ws["A19"].font = f(9, color="777777")

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = ORANGE if sheet.title == "はじめに" else "F7C08A"

wb.save(OUT)
print(OUT)
